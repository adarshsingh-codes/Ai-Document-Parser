"""
Unit tests for backend/parsers/docx_parser.py.

Fixtures are built on the fly with python-docx (already a project
dependency) rather than committing binary .docx files to the repo.
"""
import tempfile
from pathlib import Path

import pytest

from backend.parsers.base import ParsingError
from backend.parsers.docx_parser import parse_docx


def _make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> str:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)

    if table_rows:
        table = d.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row_cells = table.add_row().cells
            for i, val in enumerate(row_values):
                row_cells[i].text = val

    fd, path = tempfile.mkstemp(suffix=".docx")
    import os
    os.close(fd)
    d.save(path)
    return path


def test_parse_docx_basic_paragraphs():
    path = _make_docx(["First paragraph.", "Second paragraph."])
    try:
        result = parse_docx(path)
        assert result.file_type == "docx"
        assert "First paragraph" in result.full_text
        assert "Second paragraph" in result.full_text
        assert result.metadata["paragraph_count"] == 2
        assert result.metadata["table_count"] == 0
    finally:
        Path(path).unlink()


def test_parse_docx_skips_empty_paragraphs():
    path = _make_docx(["Real content.", "", "   ", "More real content."])
    try:
        result = parse_docx(path)
        # Only the two non-blank paragraphs should be counted.
        assert result.metadata["paragraph_count"] == 2
        assert "Real content" in result.full_text
        assert "More real content" in result.full_text
    finally:
        Path(path).unlink()


def test_parse_docx_extracts_tables():
    path = _make_docx(
        ["A document with a table below."],
        table_rows=[["Name", "Role"], ["Alice", "Engineer"], ["Bob", "Designer"]],
    )
    try:
        result = parse_docx(path)
        assert result.metadata["table_count"] == 1
        assert "[Tables]" in result.full_text
        assert "Alice" in result.full_text
        assert "Designer" in result.full_text
    finally:
        Path(path).unlink()


def test_parse_docx_missing_file_raises():
    with pytest.raises(ParsingError):
        parse_docx("/tmp/does_not_exist_ai_doc_parser.docx")


def test_parse_docx_empty_file_raises():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name  # zero bytes, not a real docx (zip) at all

    try:
        with pytest.raises(ParsingError):
            parse_docx(path)
    finally:
        Path(path).unlink()


def test_parse_docx_corrupt_file_raises():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        f.write(b"not a real zip/docx structure, just garbage bytes")
        path = f.name

    try:
        with pytest.raises(ParsingError):
            parse_docx(path)
    finally:
        Path(path).unlink()
